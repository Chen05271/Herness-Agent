"""PPT 生成工具（已废弃）— 历史 python-pptx 直出实现，Harness 请改用 ppt-master。"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from herness.agents.tools import WorkerToolError, resolve_safe_path
from herness.config import Settings

_SAFE_FILENAME_RE = re.compile(r"[^a-zA-Z0-9._-]+")


def _safe_filename(name: str) -> str:
    cleaned = _SAFE_FILENAME_RE.sub("_", name.strip()).strip("._")
    if not cleaned:
        cleaned = "presentation"
    if not cleaned.lower().endswith(".pptx"):
        cleaned += ".pptx"
    return cleaned


def _resolve_output_path(
    settings: Settings,
    *,
    task_id: str,
    filename: str,
) -> Path:
    base_dir = settings.worker_tools_output_base_dir.strip()
    if not base_dir:
        raise WorkerToolError("PPT 输出目录未配置（WORKER_TOOLS_OUTPUT_BASE_DIR）")

    task_dir = resolve_safe_path(base_dir, task_id)
    task_dir.mkdir(parents=True, exist_ok=True)
    return resolve_safe_path(str(task_dir), _safe_filename(filename))


def _parse_slides_payload(slides_json: str) -> dict[str, Any]:
    try:
        payload = json.loads(slides_json)
    except json.JSONDecodeError as exc:
        raise WorkerToolError(f"slides_json 不是合法 JSON：{exc}") from exc

    if not isinstance(payload, dict):
        raise WorkerToolError("slides_json 根节点必须是对象")
    slides = payload.get("slides")
    if not isinstance(slides, list) or not slides:
        raise WorkerToolError("slides_json 必须包含非空 slides 数组")
    return payload


def _add_title_slide(prs: Any, slide_data: dict[str, Any]) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = str(slide_data.get("title", ""))
    subtitle.text = str(slide_data.get("subtitle", ""))


def _add_content_slide(prs: Any, slide_data: dict[str, Any]) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = str(slide_data.get("title", ""))

    body = slide.placeholders[1].text_frame
    body.clear()

    bullets = slide_data.get("bullets")
    if isinstance(bullets, list) and bullets:
        for idx, item in enumerate(bullets):
            text = str(item).strip()
            if not text:
                continue
            if idx == 0:
                body.text = text
            else:
                body.add_paragraph().text = text

    table_data = slide_data.get("table")
    if isinstance(table_data, dict):
        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])
        if isinstance(headers, list) and isinstance(rows, list) and headers:
            _add_table(slide, headers, rows)

    chart_data = slide_data.get("chart")
    if isinstance(chart_data, dict):
        _add_chart(slide, chart_data)


def _add_table(slide: Any, headers: list[Any], rows: list[Any]) -> None:
    from pptx.util import Inches

    row_count = len(rows) + 1
    col_count = len(headers)
    left = Inches(0.8)
    top = Inches(2.2)
    width = Inches(8.5)
    height = Inches(0.4 * row_count)

    table_shape = slide.shapes.add_table(row_count, col_count, left, top, width, height)
    table = table_shape.table

    for col_idx, header in enumerate(headers):
        table.cell(0, col_idx).text = str(header)

    for row_idx, row in enumerate(rows, start=1):
        if not isinstance(row, list):
            continue
        for col_idx, cell in enumerate(row[:col_count]):
            table.cell(row_idx, col_idx).text = str(cell)


def _add_chart(slide: Any, chart_data: dict[str, Any]) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    chart_type = str(chart_data.get("type", "bar")).lower()
    categories = chart_data.get("categories", [])
    series_list = chart_data.get("series", [])
    if not isinstance(categories, list) or not isinstance(series_list, list):
        return

    data = CategoryChartData()
    data.categories = [str(item) for item in categories]
    for series in series_list:
        if not isinstance(series, dict):
            continue
        name = str(series.get("name", "Series"))
        values = series.get("values", [])
        if not isinstance(values, list):
            continue
        data.add_series(name, tuple(values))

    xl_type = XL_CHART_TYPE.LINE_MARKERS if chart_type == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart_frame = slide.shapes.add_chart(
        xl_type,
        Inches(0.8),
        Inches(2.0),
        Inches(8.5),
        Inches(4.0),
        data,
    )
    chart = chart_frame.chart
    chart_title = chart_data.get("title")
    if chart_title:
        chart.has_title = True
        chart.chart_title.text_frame.text = str(chart_title)


def generate_pptx(
    slides_json: str,
    *,
    filename: str,
    task_id: str,
    settings: Settings,
) -> str:
    """根据 slides JSON 生成 PPTX，返回相对输出目录的路径。"""
    if not settings.worker_tools_ppt_enabled:
        raise WorkerToolError("PPT 生成工具未启用（WORKER_TOOLS_PPT_ENABLED=false）")

    try:
        from pptx import Presentation
    except ImportError as exc:
        raise WorkerToolError(
            "缺少 python-pptx 依赖，请执行：pip install python-pptx"
        ) from exc

    payload = _parse_slides_payload(slides_json)
    output_path = _resolve_output_path(settings, task_id=task_id, filename=filename)

    prs = Presentation()
    core_props = prs.core_properties
    deck_title = str(payload.get("title", "")).strip()
    if deck_title:
        core_props.title = deck_title

    for slide_data in payload["slides"]:
        if not isinstance(slide_data, dict):
            continue
        layout = str(slide_data.get("layout", "title_content")).lower()
        if layout == "title":
            _add_title_slide(prs, slide_data)
        else:
            _add_content_slide(prs, slide_data)

    if len(prs.slides) == 0:
        raise WorkerToolError("未生成任何幻灯片，请检查 slides_json 内容")

    prs.save(str(output_path))

    base = Path(settings.worker_tools_output_base_dir.strip()).resolve()
    rel = output_path.resolve().relative_to(base)
    return str(rel).replace("\\", "/")


def resolve_artifact_path(settings: Settings, artifact_ref: str) -> Path:
    """将 artifact 相对路径解析为绝对路径并校验越界。"""
    base_dir = settings.worker_tools_output_base_dir.strip()
    if not base_dir:
        raise WorkerToolError("产物目录未配置")
    return resolve_safe_path(base_dir, artifact_ref)
